"""Excel → product_video_map.json 映射逻辑（覆盖 runtime 内旧实现）。

规则（按 Excel 整行）：
- 一行 = 一个商品（品名 + ID）
- 「视频」列一对多：逗号/中文逗号/分号/换行分隔
- 允许无后缀：`1 (1),1 (2)` → `1 (1).mp4`, `1 (2).mp4`
- 输出顺序与 Excel 行顺序一致
- 不把一行拆成多个商品；也不再按目录零散 mp4 乱配一行一个视频
"""

from __future__ import annotations

import argparse
import json
import logging
import re
from pathlib import Path
from typing import Any

import openpyxl

logger = logging.getLogger("publisher")

WORKSPACE = Path(__file__).resolve().parent.parent
DEFAULT_MAP = WORKSPACE / "data" / "product_video_map.json"

HEADER_NAME = re.compile(
    r"^(品名|商品名|名称|商品名称|product[_\s]?name|name)$", re.I
)
HEADER_ID = re.compile(r"^(id|商品\s*id|商品id|product[_\s]?id)$", re.I)
HEADER_VIDEO = re.compile(
    r"^(视频|视频名|视频名称|视频文件|视频文件名|video|videos|video[_\s]?name|files?)$",
    re.I,
)
HEADER_SKIP = re.compile(r"^(备注|说明|note|remark|skip)$", re.I)


def _clean_name(cell: Any) -> str:
    if cell is None:
        return ""
    return re.sub(r"\s+", " ", str(cell)).strip()


def _clean_product_id(cell: Any) -> str:
    if cell is None:
        return ""
    text = str(cell).strip().replace(" ", "").replace(",", "")
    # Excel 可能把长整型读成 1730....0
    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".", 1)[0]
    if re.fullmatch(r"\d+\.\d+[eE][+-]?\d+", text):
        try:
            text = str(int(float(text)))
        except Exception:
            pass
    m = re.search(r"\d{10,}", text)
    return m.group(0) if m else text


def _split_videos(cell: Any) -> list[str]:
    """拆分一对多视频：支持英文/中文逗号、分号、换行；无后缀自动补 .mp4。"""
    if cell is None:
        return []
    text = str(cell).strip()
    if not text or text.startswith("="):
        return []

    parts = re.split(r"[,，;；\n\r]+", text)
    out: list[str] = []
    seen: set[str] = set()
    for p in parts:
        p = p.strip().strip('"').strip("'").strip()
        if not p:
            continue
        # 跳过纯商品 ID
        if re.fullmatch(r"\d{15,}", p):
            continue
        if not p.lower().endswith(".mp4"):
            p = p + ".mp4"
        name = Path(p).name
        key = name.lower()
        if key in seen:
            continue
        seen.add(key)
        out.append(name)
    return out


def _detect_columns(header_row: tuple[Any, ...]) -> dict[str, int]:
    cols: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        if cell is None:
            continue
        h = str(cell).strip()
        if HEADER_NAME.match(h):
            cols["name"] = idx
        elif HEADER_ID.match(h):
            cols["id"] = idx
        elif HEADER_VIDEO.match(h):
            cols["video"] = idx
        elif HEADER_SKIP.match(h):
            cols.setdefault("skip", idx)

    if "name" not in cols or "id" not in cols:
        cols = {"name": 0, "id": 1, "video": 2 if len(header_row) > 2 else 2}
    elif "video" not in cols and len(header_row) > 2:
        cols["video"] = 2
        logger.warning(
            "未识别视频列表头，默认第 3 列作为视频列: %r", header_row[2]
        )
    return cols


def _cell(row: tuple[Any, ...], idx: int | None) -> Any:
    if idx is None or idx < 0 or idx >= len(row):
        return None
    return row[idx]


def _resolve_video_exists(videos_dir: Path, name: str) -> bool:
    p = videos_dir / name
    if p.exists():
        return True
    # 容错大小写
    for f in videos_dir.glob("*.mp4"):
        if f.name.lower() == name.lower():
            return True
    return False


def load_existing(map_file: Path) -> dict[str, dict[str, Any]]:
    if not map_file.exists():
        return {}
    try:
        data = json.loads(map_file.read_text(encoding="utf-8"))
    except Exception:
        return {}
    if not isinstance(data, list):
        return {}
    out: dict[str, dict[str, Any]] = {}
    for item in data:
        if isinstance(item, dict) and item.get("product_id"):
            out[str(item["product_id"])] = item
    return out


def guess_keywords(name: str) -> list[str]:
    return ["TikTokShop"]


def sync(excel: Path, map_file: Path, videos_dir: Path | None = None) -> list[dict[str, Any]]:
    videos_dir = Path(videos_dir) if videos_dir else WORKSPACE
    excel = Path(excel)
    map_file = Path(map_file)

    wb = openpyxl.load_workbook(excel, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows_iter = list(ws.iter_rows(values_only=True))
        if not rows_iter:
            raise RuntimeError(f"Excel 为空: {excel}")

        header = rows_iter[0]
        cols = _detect_columns(tuple(header))

        # 若首行其实是数据（表头未匹配到品名关键字），从第 0 行开始当数据
        first_name = _clean_name(_cell(header, cols.get("name")))
        start = 1
        header_name_raw = str(_cell(header, cols.get("name")) or "").strip()
        if not HEADER_NAME.match(header_name_raw):
            # 表头检测失败：若第一行像「品名+ID」数据，则纳入
            if first_name and _clean_product_id(_cell(header, cols.get("id"))):
                start = 0

        # 视频列再探测：表头未命中时，在前几行找「可拆出视频」的列
        if "video" not in cols:
            for probe in rows_iter[1:4]:
                if not probe:
                    continue
                for idx, cell in enumerate(probe):
                    if idx in (cols.get("name"), cols.get("id")):
                        continue
                    if _split_videos(cell):
                        cols["video"] = idx
                        break
                if "video" in cols:
                    break
        if "video" not in cols:
            cols["video"] = 2

        logger.info(
            "Excel 列映射 name=%s id=%s video=%s | 表头=%s",
            cols.get("name"),
            cols.get("id"),
            cols.get("video"),
            header,
        )

        data_rows: list[tuple[str, str, list[str]]] = []
        for i, row in enumerate(rows_iter[start:], start=start + 1):
            if not row:
                continue
            name = _clean_name(_cell(row, cols["name"]))
            pid = _clean_product_id(_cell(row, cols["id"]))
            if not name or not pid:
                logger.warning("跳过第 %s 行：品名或 ID 无效 -> name=%s id=%s", i, name, pid)
                continue
            vids = _split_videos(_cell(row, cols["video"]))
            missing = [v for v in vids if not _resolve_video_exists(videos_dir, v)]
            if missing:
                logger.warning(
                    "警告第 %s 行：视频文件不存在 %s（仍写入映射，enabled 将视存在情况）",
                    i,
                    missing,
                )
            data_rows.append((pid, name, vids))
    finally:
        wb.close()

    if not data_rows:
        raise RuntimeError(f"未从 Excel 解析到有效商品行: {excel}")

    existing = load_existing(map_file)
    out: list[dict[str, Any]] = []
    for pid, name, vids in data_rows:
        old = existing.get(pid) or {}
        # 只保留 Excel 本行给出的视频列表（一对多）；不再用目录剩余 mp4 自动填一行一个
        if not vids:
            # 仅当本行视频列为空时，才回退旧映射中的 videos（保持手工配置）
            old_vids = list(old.get("videos") or [])
            vids = [v for v in old_vids if isinstance(v, str) and v.strip()]
            if vids:
                logger.info("第商品 %s 视频列为空，沿用旧映射 videos=%s", pid, vids)
            else:
                logger.warning(
                    "商品 %s (%s) 无视频：Excel 视频列为空且无旧映射，enabled=false",
                    pid,
                    name,
                )

        existing_ok = [v for v in vids if _resolve_video_exists(videos_dir, v)]
        item = {
            "product_id": pid,
            "product_name": name,
            "selling_points": list(old.get("selling_points") or []),
            "videos": existing_ok if existing_ok else list(vids),
            "keywords": list(old.get("keywords") or []) or guess_keywords(name),
            "enabled": bool(existing_ok),
        }
        out.append(item)

    map_file.parent.mkdir(parents=True, exist_ok=True)
    map_file.write_text(
        json.dumps(out, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    # 旁路可读摘要
    summary = map_file.with_name("产品和视频对应.txt")
    lines = [
        f"# 来源 Excel: {excel}",
        f"# 正式映射: {map_file}",
        "# 规则: 一行=一品名+一ID+多视频(按单元格顺序)",
        "",
    ]
    for it in out:
        vids = it.get("videos") or []
        flag = "" if it.get("enabled") else "(无可用视频文件)"
        lines.append(
            f"{it['product_id']} | {it['product_name']} | videos({len(vids)})={', '.join(vids)} {flag}".rstrip()
        )
    summary.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return out


def _video_name_key(name: str) -> str:
    n = Path(str(name or "").strip()).name
    if not n.lower().endswith(".mp4"):
        n = n + ".mp4"
    return n.lower()


def remove_video_from_excel(
    excel: Path,
    product_id: str,
    video_name: str,
    *,
    videos_dir: Path | None = None,
) -> dict[str, Any]:
    """从 Excel 对应商品行的视频列中删除指定视频文本（不删整行）。

    返回: {ok, excel, product_id, video, removed, remaining, row}
    """
    excel = Path(excel)
    pid = _clean_product_id(product_id)
    target_key = _video_name_key(video_name)
    if not excel.exists():
        return {
            "ok": False,
            "excel": str(excel),
            "product_id": pid,
            "video": Path(video_name).name,
            "removed": False,
            "error": "excel_not_found",
        }

    wb = openpyxl.load_workbook(excel)
    try:
        ws = wb[wb.sheetnames[0]]
        rows_iter = list(ws.iter_rows(values_only=True))
        if not rows_iter:
            return {
                "ok": False,
                "excel": str(excel),
                "product_id": pid,
                "video": Path(video_name).name,
                "removed": False,
                "error": "excel_empty",
            }

        header = rows_iter[0]
        cols = _detect_columns(tuple(header))
        header_name_raw = str(_cell(header, cols.get("name")) or "").strip()
        start = 1
        first_name = _clean_name(_cell(header, cols.get("name")))
        if not HEADER_NAME.match(header_name_raw):
            if first_name and _clean_product_id(_cell(header, cols.get("id"))):
                start = 0
        if "video" not in cols:
            cols["video"] = 2

        id_col = int(cols["id"]) + 1  # openpyxl 1-based
        video_col = int(cols["video"]) + 1
        found_row = None
        remaining: list[str] = []
        removed = False
        original_text = ""

        for ridx in range(start + 1, ws.max_row + 1):
            row_pid = _clean_product_id(ws.cell(ridx, id_col).value)
            if row_pid != pid:
                continue
            found_row = ridx
            cell = ws.cell(ridx, video_col)
            original_text = "" if cell.value is None else str(cell.value)
            vids = _split_videos(cell.value)
            kept = [v for v in vids if _video_name_key(v) != target_key]
            removed = len(kept) < len(vids)
            remaining = kept
            # 尽量保留无后缀写法（若原单元格大多无 .mp4）
            prefer_bare = bool(original_text) and ".mp4" not in original_text.lower()
            if prefer_bare:
                write_vals = [Path(v).stem for v in kept]
            else:
                write_vals = list(kept)
            cell.value = ", ".join(write_vals) if write_vals else None
            break

        if found_row is None:
            return {
                "ok": False,
                "excel": str(excel),
                "product_id": pid,
                "video": Path(video_name).name,
                "removed": False,
                "error": "product_row_not_found",
            }

        wb.save(excel)
        logger.info(
            "Excel 已移除视频文本 product=%s video=%s row=%s remaining=%s",
            pid,
            Path(video_name).name,
            found_row,
            remaining,
        )
        return {
            "ok": True,
            "excel": str(excel),
            "product_id": pid,
            "video": Path(video_name).name,
            "removed": removed,
            "remaining": remaining,
            "row": found_row,
            "before": original_text,
        }
    finally:
        wb.close()


def rewrite_mapping_summary(excel: Path | None, map_file: Path) -> None:
    """根据当前映射重写旁路摘要「产品和视频对应.txt」。"""
    map_file = Path(map_file)
    if not map_file.exists():
        return
    try:
        data = json.loads(map_file.read_text(encoding="utf-8"))
    except Exception:
        return
    if not isinstance(data, list):
        return
    summary = map_file.with_name("产品和视频对应.txt")
    lines = [
        f"# 来源 Excel: {excel or '-'}",
        f"# 正式映射: {map_file}",
        "# 规则: 一行=一品名+一ID+多视频(按单元格顺序)",
        "",
    ]
    for it in data:
        if not isinstance(it, dict):
            continue
        vids = it.get("videos") or []
        flag = "" if it.get("enabled") else "(无可用视频文件)"
        lines.append(
            f"{it.get('product_id')} | {it.get('product_name')} | videos({len(vids)})={', '.join(vids)} {flag}".rstrip()
        )
    summary.write_text("\n".join(lines) + "\n", encoding="utf-8")


def resolve_default_excel() -> Path | None:
    desktop = Path.home() / "Desktop"
    candidates = [
        desktop / "泰国(1).xlsx",
        desktop / "泰国.xlsx",
        WORKSPACE / "泰国(1).xlsx",
    ]
    for p in candidates:
        if p.exists():
            return p
    return None


def main(argv: list[str] | None = None) -> int | None:
    parser = argparse.ArgumentParser(description="从 Excel 导入商品↔视频映射")
    parser.add_argument("--excel", type=Path, default=None, help="Excel 路径")
    parser.add_argument("--map", type=Path, default=DEFAULT_MAP, help="输出 product_video_map.json")
    parser.add_argument("--videos-dir", type=Path, default=None, help="视频目录")
    args = parser.parse_args(argv)

    excel = args.excel or resolve_default_excel()
    if not excel or not Path(excel).exists():
        print("找不到 Excel。请用 --excel 指定。")
        return 2

    videos_dir = args.videos_dir
    if videos_dir is None:
        # 兼容 UI：尽量读 config
        try:
            import config_loader

            cfg = config_loader.load_config()
            videos_dir = Path(cfg["paths"]["videos_dir"])
            if args.map == DEFAULT_MAP:
                args.map = WORKSPACE / cfg["paths"]["mapping_file"]
        except Exception:
            videos_dir = WORKSPACE

    items = sync(Path(excel), Path(args.map), Path(videos_dir))
    enabled_n = sum(1 for x in items if x.get("enabled"))
    video_n = sum(len(x.get("videos") or []) for x in items)
    print(
        f"已从 Excel 导入 {len(items)} 个商品（有效 {enabled_n}）· 映射条数(视频) {video_n} -> {args.map}"
    )
    print(f"Excel: {excel}")
    print("说明：一 ID + 一品名 + 多视频(逗号分隔，按行顺序)；发布按每条视频映射轮询。")
    for it in items:
        vids = it.get("videos") or []
        print(
            f"  - {it['product_id']} | enabled={it.get('enabled')} | videos({len(vids)})={vids}"
        )
        print(f"    {it.get('product_name')}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main() or 0)
